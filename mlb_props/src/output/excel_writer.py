"""
Professional Excel report generator for MLB Props Analyzer.

Workbook sheets:
  1. Daily_Master    – Full player table with last-N-games detail + aggregates
  2. Hits_Ranking    – Players sorted by avg Hits descending
  3. Runs_Ranking    – Players sorted by avg Runs descending
  4. RBI_Ranking     – Players sorted by avg RBI descending
  5. Data_Quality_Log – Notes and warnings collected during the pipeline
  6. Config_Metadata – Run parameters and data source info

Color conventions enforced throughout:
  • HITS  columns → gold   (#FFD966)  headers
  • RUNS  columns → green  (#92D050)  headers
  • RBI   columns → red    (#FF7070)  headers
  • Each player row → a unique light-pastel color (deterministic per player_id)
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    GradientFill,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from src.models.player_data import PlayerAnalysis
from src.output.color_manager import (
    AGG_HEADER_COLOR,
    GAME_HEADER_COLOR,
    HEADER_FONT_COLOR,
    HITS_HEADER_COLOR,
    INFO_HEADER_COLOR,
    RBI_HEADER_COLOR,
    RUNS_HEADER_COLOR,
    get_player_color,
)

logger = logging.getLogger(__name__)

# ── Thin border shared across all cells ─────────────────────────────────────
_THIN = Side(border_style="thin", color="CCCCCC")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Column width presets (characters)
_W_NAME = 26
_W_TEAM = 20
_W_POS = 9
_W_BATS = 7
_W_STATUS = 14
_W_ORDER = 9
_W_MATCHUP = 22
_W_PITCHER = 26
_W_TIME = 13
_W_HOME_AWAY = 11
_W_DATE = 13
_W_OPP = 16
_W_HIST_P = 22
_W_STAT = 8
_W_AGG_STAT = 12
_W_PCT = 10


def generate_excel_report(
    players: list[PlayerAnalysis],
    hits_ranking: list[dict[str, Any]],
    runs_ranking: list[dict[str, Any]],
    rbi_ranking: list[dict[str, Any]],
    run_date: str,
    output_dir: str = "output",
    max_lookback: int = 10,
) -> str:
    """
    Build and save the full Excel workbook.

    Returns the absolute path of the saved file.
    """
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    filename = f"MLB_Props_{run_date}.xlsx"
    filepath = str(Path(output_dir) / filename)

    wb = Workbook()

    # ── Sheet 1: Daily Master ────────────────────────────────────────────────
    ws_master = wb.active
    ws_master.title = "Daily_Master"
    _write_master_sheet(ws_master, players, max_lookback)

    # ── Sheets 2–4: Rankings ────────────────────────────────────────────────
    ws_hits = wb.create_sheet("Hits_Ranking")
    _write_ranking_sheet(ws_hits, hits_ranking, "hits", "Hits")

    ws_runs = wb.create_sheet("Runs_Ranking")
    _write_ranking_sheet(ws_runs, runs_ranking, "runs", "Runs")

    ws_rbi = wb.create_sheet("RBI_Ranking")
    _write_ranking_sheet(ws_rbi, rbi_ranking, "rbi", "RBI")

    # ── Sheet 5: Data Quality Log ────────────────────────────────────────────
    ws_log = wb.create_sheet("Data_Quality_Log")
    _write_quality_log(ws_log, players)

    # ── Sheet 6: Config / Metadata ───────────────────────────────────────────
    ws_meta = wb.create_sheet("Config_Metadata")
    _write_metadata(ws_meta, run_date, len(players), max_lookback)

    wb.save(filepath)
    logger.info("Excel report saved: %s", filepath)
    return filepath


# ══════════════════════════════════════════════════════════════════════════════
# DAILY MASTER SHEET
# ══════════════════════════════════════════════════════════════════════════════

def _write_master_sheet(
    ws: Worksheet,
    players: list[PlayerAnalysis],
    max_lookback: int,
) -> None:
    """Write the full player table to the Daily_Master sheet."""

    # ── Build column definitions ─────────────────────────────────────────────
    col_defs = _build_master_columns(max_lookback)
    headers = [c["label"] for c in col_defs]
    widths = [c["width"] for c in col_defs]
    categories = [c["cat"] for c in col_defs]  # "info"|"game"|"hist_info"|"H"|"R"|"RBI"|"agg_*"

    # Title row
    ws.row_dimensions[1].height = 20
    title_cell = ws.cell(row=1, column=1, value=f"MLB Props Analyzer — {_today_label()}")
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F3864")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))

    # Header row (row 2)
    ws.row_dimensions[2].height = 32
    for col_idx, (header, cat, width) in enumerate(zip(headers, categories, widths), start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = Font(name="Calibri", bold=True, size=9, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=_header_color(cat))
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = _CELL_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze title + header rows and player-identity columns
    ws.freeze_panes = "E3"

    # Data rows
    for row_idx, player in enumerate(players, start=3):
        player_color = get_player_color(player.player_id)
        row_fill = PatternFill("solid", fgColor=player_color)
        row_data = _player_to_master_row(player, max_lookback)

        ws.row_dimensions[row_idx].height = 16

        for col_idx, (value, cat) in enumerate(zip(row_data, categories), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.font = Font(name="Calibri", size=9)
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(horizontal=_cell_align(cat), vertical="center")

            # Numeric formatting
            if cat in ("H", "R", "RBI") and isinstance(value, int):
                cell.number_format = "0"
            elif cat.startswith("agg_avg"):
                cell.number_format = "0.000"
            elif cat.startswith("agg_pct"):
                cell.number_format = "0.0"

    # Auto-filter on header row
    ws.auto_filter.ref = (
        f"A2:{get_column_letter(len(headers))}{len(players) + 2}"
    )

    logger.debug("Master sheet: %d rows, %d columns", len(players), len(headers))


def _build_master_columns(max_lookback: int) -> list[dict]:
    """Return ordered column definitions: label, category key, width."""
    cols: list[dict] = [
        {"label": "Player Name",     "cat": "info", "width": _W_NAME},
        {"label": "Team",            "cat": "info", "width": _W_TEAM},
        {"label": "Position",        "cat": "info", "width": _W_POS},
        {"label": "Bats",            "cat": "info", "width": _W_BATS},
        {"label": "Lineup Status",   "cat": "info", "width": _W_STATUS},
        {"label": "Bat Order",       "cat": "info", "width": _W_ORDER},
        {"label": "Today Matchup",   "cat": "game", "width": _W_MATCHUP},
        {"label": "Today Pitcher",   "cat": "game", "width": _W_PITCHER},
        {"label": "Game Time",       "cat": "game", "width": _W_TIME},
        {"label": "Home/Away",       "cat": "game", "width": _W_HOME_AWAY},
    ]

    for i in range(1, max_lookback + 1):
        cols += [
            {"label": f"G{i} Date",    "cat": "hist_info", "width": _W_DATE},
            {"label": f"G{i} Opp",     "cat": "hist_info", "width": _W_OPP},
            {"label": f"G{i} Pitcher", "cat": "hist_info", "width": _W_HIST_P},
            {"label": f"G{i} H",       "cat": "H",         "width": _W_STAT},
            {"label": f"G{i} R",       "cat": "R",         "width": _W_STAT},
            {"label": f"G{i} RBI",     "cat": "RBI",       "width": _W_STAT},
        ]

    cols += [
        {"label": "Total H",       "cat": "agg_h",    "width": _W_AGG_STAT},
        {"label": "Total R",       "cat": "agg_r",    "width": _W_AGG_STAT},
        {"label": "Total RBI",     "cat": "agg_rbi",  "width": _W_AGG_STAT},
        {"label": "Avg H",         "cat": "agg_avg_h",   "width": _W_AGG_STAT},
        {"label": "Avg R",         "cat": "agg_avg_r",   "width": _W_AGG_STAT},
        {"label": "Avg RBI",       "cat": "agg_avg_rbi", "width": _W_AGG_STAT},
        {"label": "G 1+H",         "cat": "agg_gh",   "width": _W_STAT},
        {"label": "G 1+R",         "cat": "agg_gr",   "width": _W_STAT},
        {"label": "G 1+RBI",       "cat": "agg_grbi", "width": _W_STAT},
        {"label": "% 1+H",         "cat": "agg_pct_h",   "width": _W_PCT},
        {"label": "% 1+R",         "cat": "agg_pct_r",   "width": _W_PCT},
        {"label": "% 1+RBI",       "cat": "agg_pct_rbi", "width": _W_PCT},
        {"label": "Games Used",    "cat": "agg_meta", "width": _W_STAT},
        {"label": "Form Trend",    "cat": "agg_meta", "width": _W_STATUS},
    ]

    return cols


def _player_to_master_row(player: PlayerAnalysis, max_lookback: int) -> list:
    """Flatten a PlayerAnalysis into an ordered list of cell values."""
    row: list = [
        player.name,
        player.team,
        player.position,
        player.bats,
        player.lineup_status,
        player.batting_order if player.batting_order else "N/A",
        player.today_game.matchup_str,
        player.today_game.opposing_pitcher,
        player.today_game.game_time_local,
        player.today_game.venue_label,
    ]

    for i in range(max_lookback):
        if i < len(player.last_n_games):
            g = player.last_n_games[i]
            row += [g.date, g.opponent, g.opposing_pitcher, g.hits, g.runs, g.rbi]
        else:
            row += ["—", "—", "—", "—", "—", "—"]

    row += [
        player.total_hits,
        player.total_runs,
        player.total_rbi,
        player.avg_hits,
        player.avg_runs,
        player.avg_rbi,
        player.games_with_hit,
        player.games_with_run,
        player.games_with_rbi,
        player.pct_hit,
        player.pct_run,
        player.pct_rbi,
        player.games_used,
        player.form_trend or "N/A",
    ]

    return row


# ══════════════════════════════════════════════════════════════════════════════
# RANKING SHEETS
# ══════════════════════════════════════════════════════════════════════════════

def _write_ranking_sheet(
    ws: Worksheet,
    ranking: list[dict[str, Any]],
    metric: str,   # "hits" | "runs" | "rbi"
    label: str,    # "Hits" | "Runs" | "RBI"
) -> None:
    if not ranking:
        ws.cell(row=1, column=1, value="No data available")
        return

    # Colour for this metric
    metric_color = {"hits": HITS_HEADER_COLOR, "runs": RUNS_HEADER_COLOR, "rbi": RBI_HEADER_COLOR}[metric]

    headers = list(ranking[0].keys())
    # Remove internal Player_ID from display
    display_headers = [h for h in headers if h != "Player_ID"]

    # Title row
    ws.row_dimensions[1].height = 20
    title = ws.cell(row=1, column=1, value=f"MLB Props — {label} Ranking")
    title.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F3864")
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(display_headers))

    # Header row
    ws.row_dimensions[2].height = 28
    col_widths = {
        "Rank": 7, "Player": 26, "Team": 20, "Today_Matchup": 20,
        "Today_Pitcher": 26, "Lineup_Status": 14, "Batting_Order": 10,
        "Games_Used": 10, "Form_Trend": 12,
    }

    for col_idx, h in enumerate(display_headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        # Metric-specific columns get metric color; others get neutral
        is_metric_col = (
            metric.capitalize() in h
            or label in h
            or "Avg" in h
            or "Total" in h
            or "Games_1" in h
            or "Pct" in h
        )
        header_fill = metric_color if is_metric_col else INFO_HEADER_COLOR
        cell.font = Font(name="Calibri", bold=True, size=9, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _CELL_BORDER

        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_widths.get(h, 14)

    # Freeze header rows
    ws.freeze_panes = "A3"

    # Data rows
    for row_idx, row_data in enumerate(ranking, start=3):
        player_id = row_data.get("Player_ID", 0)
        player_color = get_player_color(player_id)
        row_fill = PatternFill("solid", fgColor=player_color)

        ws.row_dimensions[row_idx].height = 16

        for col_idx, h in enumerate(display_headers, start=1):
            value = row_data.get(h, "N/A")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.font = Font(name="Calibri", size=9)
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(
                horizontal="right" if isinstance(value, (int, float)) else "left",
                vertical="center",
            )
            if isinstance(value, float):
                cell.number_format = "0.000" if "Avg" in h else "0.0"

    ws.auto_filter.ref = f"A2:{get_column_letter(len(display_headers))}{len(ranking) + 2}"


# ══════════════════════════════════════════════════════════════════════════════
# DATA QUALITY LOG SHEET
# ══════════════════════════════════════════════════════════════════════════════

def _write_quality_log(ws: Worksheet, players: list[PlayerAnalysis]) -> None:
    headers = ["Player", "Team", "Games_Used", "Lineup_Status", "Notes"]
    _write_simple_header(ws, headers, "Data Quality Log", "44546A")

    for row_idx, p in enumerate(players, start=3):
        notes = "; ".join(p.data_quality_notes) if p.data_quality_notes else "OK"
        row = [p.name, p.team, p.games_used, p.lineup_status, notes]
        for col_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=9)
            cell.border = _CELL_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx, w in enumerate([26, 20, 10, 14, 80], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.freeze_panes = "A3"


# ══════════════════════════════════════════════════════════════════════════════
# CONFIG / METADATA SHEET
# ══════════════════════════════════════════════════════════════════════════════

def _write_metadata(
    ws: Worksheet,
    run_date: str,
    player_count: int,
    max_lookback: int,
) -> None:
    meta_rows = [
        ("Run Date", run_date),
        ("Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Players Analyzed", player_count),
        ("Games Lookback", max_lookback),
        ("Data Source", "MLB Stats API — statsapi.mlb.com (public, no key required)"),
        ("Version", "1.0.0"),
        ("Notes", "Opposing pitchers resolved via boxscore API (cached per game_pk)."),
        (
            "Future Extensions",
            "Confidence score, home/away splits, vs-hand splits, parlay optimizer.",
        ),
    ]

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 65

    title = ws.cell(row=1, column=1, value="MLB Props Analyzer — Run Metadata")
    title.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F3864")
    ws.merge_cells("A1:B1")

    for row_idx, (key, val) in enumerate(meta_rows, start=2):
        k_cell = ws.cell(row=row_idx, column=1, value=key)
        k_cell.font = Font(name="Calibri", bold=True, size=10)
        k_cell.fill = PatternFill("solid", fgColor="D9E1F2")

        v_cell = ws.cell(row=row_idx, column=2, value=val)
        v_cell.font = Font(name="Calibri", size=10)


# ══════════════════════════════════════════════════════════════════════════════
# HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def _header_color(cat: str) -> str:
    """Map a column category key to its header background color."""
    mapping = {
        "info":      INFO_HEADER_COLOR,
        "game":      GAME_HEADER_COLOR,
        "hist_info": "808080",    # Mid-gray for historical context cols
        "H":         HITS_HEADER_COLOR,
        "R":         RUNS_HEADER_COLOR,
        "RBI":       RBI_HEADER_COLOR,
        "agg_h":     HITS_HEADER_COLOR,
        "agg_r":     RUNS_HEADER_COLOR,
        "agg_rbi":   RBI_HEADER_COLOR,
        "agg_avg_h": HITS_HEADER_COLOR,
        "agg_avg_r": RUNS_HEADER_COLOR,
        "agg_avg_rbi": RBI_HEADER_COLOR,
        "agg_gh":    HITS_HEADER_COLOR,
        "agg_gr":    RUNS_HEADER_COLOR,
        "agg_grbi":  RBI_HEADER_COLOR,
        "agg_pct_h": HITS_HEADER_COLOR,
        "agg_pct_r": RUNS_HEADER_COLOR,
        "agg_pct_rbi": RBI_HEADER_COLOR,
        "agg_meta":  AGG_HEADER_COLOR,
    }
    return mapping.get(cat, INFO_HEADER_COLOR)


def _cell_align(cat: str) -> str:
    """Horizontal alignment for data cells based on column category."""
    numeric_cats = {"H", "R", "RBI", "agg_h", "agg_r", "agg_rbi",
                    "agg_avg_h", "agg_avg_r", "agg_avg_rbi",
                    "agg_gh", "agg_gr", "agg_grbi",
                    "agg_pct_h", "agg_pct_r", "agg_pct_rbi"}
    return "center" if cat in numeric_cats else "left"


def _write_simple_header(
    ws: Worksheet,
    headers: list[str],
    title: str,
    color: str,
) -> None:
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(name="Calibri", bold=True, size=13, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F3864")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = Font(name="Calibri", bold=True, size=9, color=HEADER_FONT_COLOR)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _CELL_BORDER

    ws.row_dimensions[2].height = 24


def _today_label() -> str:
    return datetime.now().strftime("%B %d, %Y")
